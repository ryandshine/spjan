#!/usr/bin/env python3
"""Membuat packages/shared/src/data/sbm-2026.ts dari berkas 'SPJ Otomatis.xlsx' (prototipe spreadsheet yang memuat
sheet tabel SBM: Uang Harian, Uang Representasi, Penginapan, Transport ..., Tiket Pesawat PP, Sewa Kendaraan).

Pemakaian: python3 scripts/build-sbm-fixture.py "<path ke xlsx>" [keluaran.ts]
Tabel-tabel itu diekstrak dari PMK 32 Tahun 2025 (Lampiran I dan II). Berkas xlsx tidak disimpan di repo.
"""
import json
import re
import sys

import openpyxl

SRC = sys.argv[1]
OUT = sys.argv[2] if len(sys.argv) > 2 else "packages/shared/src/data/sbm-2026.ts"
wb = openpyxl.load_workbook(SRC, data_only=True)


def norm(s):
    return re.sub(r"\s+", " ", str(s).strip()).upper()


def rows(sheet, first, last, ncols):
    ws = wb[sheet]
    for r in range(first, last + 1):
        yield [ws.cell(r, c).value for c in range(1, ncols + 1)]


uang_harian = {norm(r[1]): {"luarKota": r[3], "dalamKota8Jam": r[4], "diklat": r[5]} for r in rows("Uang Harian", 6, 43, 6)}
representasi = [{"uraian": r[1], "luarKota": r[3], "dalamKota": r[4]} for r in rows("Uang Representasi", 6, 8, 5)]
penginapan = {norm(r[1]): [r[3], r[4], r[5], r[6]] for r in rows("Penginapan", 6, 43, 7)}
terminal = {norm(r[1]): r[3] for r in rows("Transport Bandara-Terminal", 6, 39, 4)}
tiket = [{"asal": norm(r[1]), "tujuan": norm(r[2]), "bisnis": r[3], "ekonomi": r[4]} for r in rows("Tiket Pesawat PP", 6, 321, 5)]
ibukota = [
    {"provinsi": norm(r[1]), "ibukota": r[2], "kabKota": r[3], "besaran": r[5]}
    for r in rows("Transport Ibukota-Kab", 6, 366, 6)
]
jakarta = [{"kabKota": r[2], "besaran": r[4]} for r in rows("Transport Jakarta-Sekitar", 6, 14, 5)]
kegiatan_pp = wb["Transport Jakarta-Sekitar"].cell(19, 5).value
sewa = {norm(r[1]): {"roda4": r[3], "roda6BusSedang": r[4], "busBesar": r[5]} for r in rows("Sewa Kendaraan", 6, 43, 6)}

expected = {
    "uangHarian": (uang_harian, 38), "penginapan": (penginapan, 38), "terminal": (terminal, 34), "tiket": (tiket, 316),
    "ibukota": (ibukota, 361), "jakarta": (jakarta, 9), "sewa": (sewa, 38), "representasi": (representasi, 3),
}
for nama, (data, n) in expected.items():
    assert len(data) == n, f"{nama}: {len(data)} baris, seharusnya {n}"
assert kegiatan_pp == 170000, kegiatan_pp

sbm = {
    "tahunAnggaran": 2026,
    "uangHarian": uang_harian,
    "representasi": representasi,
    "penginapan": penginapan,
    "terminal": terminal,
    "tiketPp": tiket,
    "transportIbukota": ibukota,
    "transportJakarta": jakarta,
    "transportKegiatanPp": kegiatan_pp,
    "sewaKendaraan": sewa,
    "statusKonfigurasi": {
        "PEJABAT_NEGARA_WAMEN": {"kolomHotel": 1, "barisRepresentasi": 1, "kelasTiket": "Bisnis"},
        "ESELON_I": {"kolomHotel": 1, "barisRepresentasi": 2, "kelasTiket": "Bisnis"},
        "ESELON_II": {"kolomHotel": 2, "barisRepresentasi": 3, "kelasTiket": "Ekonomi"},
        "ESELON_III_GOL_IV": {"kolomHotel": 3, "barisRepresentasi": 0, "kelasTiket": "Ekonomi"},
        "ESELON_IV_GOL_III_II_I": {"kolomHotel": 4, "barisRepresentasi": 0, "kelasTiket": "Ekonomi"},
    },
}

with open(OUT, "w", encoding="utf-8") as f:
    f.write("// Dihasilkan oleh scripts/build-sbm-fixture.py dari PMK 32 Tahun 2025. Jangan diedit manual.\n")
    f.write('import type { Sbm } from "../types.js";\n\n')
    f.write("export const SBM_2026: Sbm = ")
    f.write(json.dumps(sbm, ensure_ascii=False, indent=1))
    f.write(";\n")
print("ditulis", OUT)
